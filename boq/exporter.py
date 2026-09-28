"""
BOQ export (Section 4.2's "Export BOQ to Excel/PDF"), in the "standard
layout with bill totals and a summary page" the spec asks for:

  - Excel: one "BOQ" sheet with every bill's items (a bill-title row,
    its items, then a bold subtotal row), and a "Summary" sheet listing
    each bill's total plus the grand total — the same figures
    `Bill.total` / `BOQ.grand_total` already compute, written out
    rather than recalculated by the spreadsheet.
  - PDF: the same data rendered from an HTML template via WeasyPrint.

Both read-only: neither function writes to the database, and both are
reachable by anyone who can *view* the BOQ (not just edit it).
"""

import io

from django.template.loader import render_to_string
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

_COLUMN_HEADERS = ["Item Ref", "Description", "Unit", "Quantity", "Rate", "Amount"]
_COLUMN_WIDTHS = [12, 50, 10, 14, 14, 16]


def _write_boq_sheet(ws, boq):
    ws.title = "BOQ"
    for col, (header, width) in enumerate(zip(_COLUMN_HEADERS, _COLUMN_WIDTHS), start=1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = Font(bold=True)
        ws.column_dimensions[get_column_letter(col)].width = width

    row = 2
    for bill in boq.bills.all():
        title_cell = ws.cell(row=row, column=1, value=f"Bill {bill.number}: {bill.title}")
        title_cell.font = Font(bold=True)
        row += 1

        for item in bill.items.all():
            if item.item_type == item.TYPE_HEADING:
                cell = ws.cell(row=row, column=2, value=item.description)
                cell.font = Font(bold=True, italic=True)
                ws.cell(row=row, column=1, value=item.item_reference)
            else:
                ws.cell(row=row, column=1, value=item.item_reference)
                ws.cell(row=row, column=2, value=item.description)
                ws.cell(row=row, column=3, value=item.unit.code if item.unit else "")
                ws.cell(row=row, column=4, value=float(item.quantity) if item.quantity is not None else None)
                ws.cell(row=row, column=5, value=float(item.rate) if item.rate is not None else None)
                ws.cell(row=row, column=6, value=float(item.amount) if item.amount is not None else None)
            row += 1

        subtotal_label = ws.cell(row=row, column=5, value=f"Bill {bill.number} total")
        subtotal_label.font = Font(bold=True)
        subtotal_label.alignment = Alignment(horizontal="right")
        subtotal_value = ws.cell(row=row, column=6, value=float(bill.total))
        subtotal_value.font = Font(bold=True)
        row += 2  # a blank row between bills


def _write_summary_sheet(ws, boq):
    ws.title = "Summary"
    for col, header in enumerate(["Bill Number", "Bill Title", "Total"], start=1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = Font(bold=True)
    ws.column_dimensions["A"].width = 12
    ws.column_dimensions["B"].width = 50
    ws.column_dimensions["C"].width = 16

    row = 2
    for bill in boq.bills.all():
        ws.cell(row=row, column=1, value=bill.number)
        ws.cell(row=row, column=2, value=bill.title)
        ws.cell(row=row, column=3, value=float(bill.total))
        row += 1

    grand_label = ws.cell(row=row + 1, column=2, value="Grand total")
    grand_label.font = Font(bold=True)
    grand_value = ws.cell(row=row + 1, column=3, value=float(boq.grand_total))
    grand_value.font = Font(bold=True)


def build_workbook(boq):
    """An openpyxl Workbook with a BOQ sheet and a Summary sheet."""
    wb = Workbook()
    _write_boq_sheet(wb.active, boq)
    _write_summary_sheet(wb.create_sheet(), boq)
    return wb


def export_xlsx_bytes(boq):
    """The workbook above, as .xlsx bytes ready for an HTTP response."""
    wb = build_workbook(boq)
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def export_pdf_bytes(boq):
    """The BOQ rendered to PDF via WeasyPrint, same figures as the Excel export."""
    from weasyprint import HTML  # imported lazily: only needed for this one path

    html = render_to_string(
        "boq/export_pdf.html",
        {
            "boq": boq,
            "bills": boq.bills.all(),
        },
    )
    return HTML(string=html).write_pdf()


TEMPLATE_HEADERS = [
    "Bill Number",
    "Bill Title",
    "Item Reference",
    "Description",
    "Unit",
    "Quantity",
    "Rate",
    "Section",
]

TEMPLATE_EXAMPLE_ROWS = [
    [4, "Earthworks", "", "Site clearance and earthworks", "", "", "", ""],
    [4, "Earthworks", "4.01", "Clear site vegetation", "m2", 1500, 2.50, ""],
    [4, "Earthworks", "4.02", "Excavate to reduce level", "m3", 1250.500, 185.00, ""],
    [4, "Earthworks", "4.03", "Provisional sum for unforeseen ground conditions", "sum", "", 50000.00, ""],
]


def build_import_template_bytes():
    """
    A small .xlsx with the flat-layout header row `import.py` expects,
    plus a few example rows (a heading, a measured item, and a lump
    sum) so a new user can see the shape without guessing at it.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "BOQ Import"
    for col, header in enumerate(TEMPLATE_HEADERS, start=1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = Font(bold=True)
        ws.column_dimensions[get_column_letter(col)].width = 20
    for row_offset, example_row in enumerate(TEMPLATE_EXAMPLE_ROWS, start=2):
        for col, value in enumerate(example_row, start=1):
            ws.cell(row=row_offset, column=col, value=value if value != "" else None)
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
