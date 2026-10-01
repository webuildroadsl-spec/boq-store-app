"""
BOQ Excel import (Section 4.2's "Import BOQ from Excel").

Scope, agreed with the user before building this:
  - Expects a flat layout: one row per item, with a bill number and
    bill title repeated on every row (matching the downloadable
    template), not a sheet with merged bill-header rows. Real BOQ
    layouts vary a lot (see the spec's own open question about this),
    so this targets the template rather than guessing at an arbitrary
    existing layout.
  - Item type isn't a mapped column. It's inferred: a row with no
    quantity and no rate is a Heading; a row whose unit is "sum" is a
    Lump Sum; everything else is Measured. Provisional Sum / Prime
    Cost / Daywork rows import as one of those two and can be
    reclassified afterward in the manual-entry grid.

Column mapping is by header name (case/spacing-insensitive, matched
against a short alias list), with a manual fallback the person can
adjust if their headers don't match.
"""

import re
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation

import openpyxl

from core.models import UnitOfMeasure
from .models import BOQItem

REQUIRED_FIELDS = ["bill_number", "bill_title", "item_reference", "description"]
OPTIONAL_FIELDS = ["unit", "quantity", "rate", "section_code"]
ALL_FIELDS = REQUIRED_FIELDS + OPTIONAL_FIELDS

FIELD_LABELS = {
    "bill_number": "Bill number",
    "bill_title": "Bill title",
    "item_reference": "Item reference",
    "description": "Description",
    "unit": "Unit",
    "quantity": "Quantity",
    "rate": "Rate",
    "section_code": "Section",
}

_ALIASES = {
    "bill_number": ["bill number", "bill no", "bill", "bill #"],
    "bill_title": ["bill title", "bill name", "title"],
    "item_reference": ["item reference", "item ref", "ref", "reference", "item no", "item number"],
    "description": ["description", "desc"],
    "unit": ["unit", "uom", "unit of measure"],
    "quantity": ["quantity", "qty"],
    "rate": ["rate", "unit rate", "price"],
    "section_code": ["section", "section code", "chainage"],
}


# Other spellings of the unit codes seen in real BOQs, keyed by
# _unit_key() of the spelling. An exact code match is tried first.
_UNIT_ALIASES = {
    "m3": "m³", "cum": "m³", "cubm": "m³",
    "m2": "m²", "sqm": "m²",
    "lm": "m",
    "no": "nr", "nos": "nr", "number": "nr", "each": "nr", "ea": "nr", "pcs": "nr",
    "ls": "sum", "lumpsum": "sum",
    "tonne": "t", "tonnes": "t", "ton": "t", "tons": "t",
    "kgs": "kg",
    "ltr": "L", "litre": "L", "litres": "L", "liter": "L",
    "hrs": "hr", "hour": "hr", "hours": "hr",
    "days": "day",
    "hectare": "ha", "hectares": "ha",
    "m3km": "m³·km", "m³km": "m³·km",
    "tkm": "t·km",
}


def _unit_key(text):
    """'L.S.' -> 'ls', 'm3*km' -> 'm3km', 'No.' -> 'no', 'm³·km' -> 'm³km'."""
    return re.sub(r"[\s.*·×/-]+", "", str(text).strip().lower())


def find_unit(raw, units_by_code):
    """The UnitOfMeasure for a unit cell, or None if it isn't recognised."""
    unit = units_by_code.get(str(raw).strip().lower())
    if unit is not None:
        return unit
    code = _UNIT_ALIASES.get(_unit_key(raw))
    return units_by_code.get(code.lower()) if code else None


def _normalize(text):
    return re.sub(r"[^a-z0-9]+", " ", (text or "").strip().lower()).strip()


def read_headers(path):
    """The first row of the uploaded file, as plain strings."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb.active
        row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), ())
        return ["" if cell is None else str(cell).strip() for cell in row]
    finally:
        wb.close()


def auto_detect_mapping(headers):
    """{field: column_index} for every field whose header we could match."""
    normalized = [_normalize(h) for h in headers]
    mapping = {}
    for field, aliases in _ALIASES.items():
        for i, header in enumerate(normalized):
            if header in aliases:
                mapping[field] = i
                break
    return mapping


def _cell(row, mapping, field):
    index = mapping.get(field)
    if index is None or index >= len(row):
        return None
    value = row[index]
    if isinstance(value, str):
        value = value.strip()
    return value if value not in ("",) else None


def _to_decimal(value):
    if value is None:
        return None, None
    try:
        return Decimal(str(value)), None
    except InvalidOperation:
        return None, f"'{value}' is not a valid number."


def parse_rows(path, mapping, project, boq):
    """
    Reads every data row (from row 2) and returns a list of dicts:
    {row_number, bill_number, bill_title, item_reference, description,
     unit_code, unit_id, quantity, rate, item_type, amount, section_code,
     section_id, errors: [str, ...]}

    Nothing is written to the database — this only reads and validates,
    so the same function drives both the preview screen and the
    server-side re-check the confirm step does before saving.
    """
    units_by_code = {u.code.lower(): u for u in UnitOfMeasure.objects.all()}
    sections_by_code = {s.code.lower(): s for s in project.sections.all()}
    existing_refs = set(
        BOQItem.objects.filter(bill__boq=boq).values_list("item_reference", flat=True)
    )

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb.active
        rows = []
        refs_seen_in_file = {}
        bill_titles_seen = {}

        for row_number, raw_row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            bill_number_raw = _cell(raw_row, mapping, "bill_number")
            item_reference = _cell(raw_row, mapping, "item_reference")
            description = _cell(raw_row, mapping, "description")
            if bill_number_raw is None and item_reference is None and description is None:
                continue  # a fully blank row: not part of the BOQ.

            errors = []

            bill_number = None
            if bill_number_raw is None:
                errors.append("Bill number is required.")
            else:
                try:
                    bill_number = int(bill_number_raw)
                except (TypeError, ValueError):
                    errors.append(f"Bill number '{bill_number_raw}' is not a whole number.")

            bill_title_raw = _cell(raw_row, mapping, "bill_title")
            if bill_number is not None:
                if bill_number not in bill_titles_seen:
                    if not bill_title_raw:
                        errors.append("Bill title is required on this bill's first row.")
                    else:
                        bill_titles_seen[bill_number] = str(bill_title_raw)
                bill_title = bill_titles_seen.get(bill_number, bill_title_raw or "")
            else:
                bill_title = bill_title_raw or ""

            if not item_reference:
                errors.append("Item reference is required.")
            else:
                item_reference = str(item_reference)
                if item_reference in existing_refs:
                    errors.append(
                        f"Item reference '{item_reference}' already exists in this BOQ version."
                    )
                elif item_reference in refs_seen_in_file:
                    errors.append(
                        f"Item reference '{item_reference}' is duplicated on row "
                        f"{refs_seen_in_file[item_reference]}."
                    )
                else:
                    refs_seen_in_file[item_reference] = row_number

            if not description:
                errors.append("Description is required.")

            unit_raw = _cell(raw_row, mapping, "unit")
            quantity_raw = _cell(raw_row, mapping, "quantity")
            rate_raw = _cell(raw_row, mapping, "rate")

            quantity, quantity_error = _to_decimal(quantity_raw)
            if quantity_error:
                errors.append(quantity_error)
            rate, rate_error = _to_decimal(rate_raw)
            if rate_error:
                errors.append(rate_error)

            is_heading = quantity is None and rate is None and not quantity_error and not rate_error

            unit = None
            item_type = None
            amount = None

            if is_heading:
                item_type = BOQItem.TYPE_HEADING
                quantity = None
                rate = None
            else:
                if not unit_raw:
                    errors.append("Unit is required for a non-heading row.")
                else:
                    unit = find_unit(unit_raw, units_by_code)
                    if unit is None:
                        errors.append(f"Unit '{unit_raw}' is not in the unit-of-measure list.")

                is_lump_sum = unit is not None and unit.code.lower() == "sum"
                if is_lump_sum:
                    # Rule 2: a lump sum is entered as one figure — the
                    # quantity/rate-required checks below don't apply to it.
                    item_type = BOQItem.TYPE_LUMP_SUM
                    quantity = Decimal("1")
                    if rate is None:
                        errors.append("Rate is required for a lump sum item.")
                    else:
                        amount = rate.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
                else:
                    if quantity is None:
                        errors.append("Quantity is required unless both quantity and rate are blank (a heading).")
                    if rate is None:
                        errors.append("Rate is required unless both quantity and rate are blank (a heading).")
                    if quantity is not None and rate is not None:
                        item_type = BOQItem.TYPE_MEASURED
                        amount = (quantity * rate).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)

            section_id = None
            section_code = _cell(raw_row, mapping, "section_code")
            if section_code:
                section = sections_by_code.get(str(section_code).lower())
                if section is None:
                    errors.append(f"Section '{section_code}' was not found on this project.")
                else:
                    section_id = section.pk

            rows.append(
                {
                    "row_number": row_number,
                    "bill_number": bill_number,
                    "bill_title": bill_title,
                    "item_reference": item_reference or "",
                    "description": description or "",
                    "unit_code": unit.code if unit else (unit_raw or ""),
                    "unit_id": unit.pk if unit else None,
                    "quantity": str(quantity) if quantity is not None else "",
                    "rate": str(rate) if rate is not None else "",
                    "item_type": item_type,
                    "amount": str(amount) if amount is not None else "",
                    "section_code": section_code or "",
                    "section_id": section_id,
                    "errors": errors,
                }
            )
        return rows
    finally:
        wb.close()
